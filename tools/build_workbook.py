"""Build FX_Hedging_Tool.xlsm: sheets, live formulas, names, buttons, charts + VBA project.

    python tools/build_workbook.py                       # formulas + empty Monte Carlo outputs
    python tools/build_workbook.py --prefill run.xlsx    # cached values from a recalculated run

Design: everything that has a closed form is a worksheet formula (Schedule sheet,
Pricer mid values, Scenarios grid, Pitch). VBA (vba/*.bas) provides the market
worksheet functions (FX_FORWARD, FX_VOL, ...), the zero-cost strike solver and the
Monte Carlo. This script only lays the workbook out so it can be rebuilt reproducibly.
"""
from __future__ import annotations

import argparse
import math
import os
import sys

import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_vba_project import build as build_vba  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEETS = ["Pricer", "Pitch", "Schedule", "Scenarios", "Distribution", "Market", "Notes"]
STRATS = ["Unhedged", "Forward", "Vanilla option", "Zero-cost collar", "Participating forward", "TARF"]
COLORS = ["#8a8985", "#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"]
NFIX = 24          # schedule rows (max number of monthly flows)
SROW = 2           # first schedule data row (0-based) -> Excel row 3
NSCEN = 61

WORKBOOK_CODE = '''
Private Sub Workbook_Open()
    ' worksheet functions need the VBA project: recalculate once macros are enabled
    On Error Resume Next
    Application.CalculateFull
End Sub
'''

MARKET = [  # tenor, years, USD %, EUR %, ATM %, RR %, BF %
    ("1W", 7 / 365, 3.85, 1.95, 6.90, 0.15, 0.12),
    ("1M", 1 / 12, 3.80, 1.95, 7.10, 0.20, 0.14),
    ("2M", 2 / 12, 3.75, 1.96, 7.25, 0.25, 0.16),
    ("3M", 0.25, 3.70, 1.97, 7.40, 0.30, 0.18),
    ("6M", 0.5, 3.55, 2.00, 7.70, 0.40, 0.21),
    ("9M", 0.75, 3.45, 2.03, 7.90, 0.45, 0.23),
    ("1Y", 1.0, 3.35, 2.05, 8.05, 0.50, 0.25),
    ("18M", 1.5, 3.25, 2.10, 8.20, 0.50, 0.27),
    ("2Y", 2.0, 3.20, 2.15, 8.30, 0.50, 0.28),
]


class Cache:
    """Values of a recalculated copy of the workbook (LibreOffice run), used as cached results."""

    def __init__(self, path):
        self.wb = None
        if path:
            import openpyxl
            self.wb = openpyxl.load_workbook(path, data_only=True)

    def __call__(self, sheet, row, col, default=0):
        if self.wb is None:
            return default
        v = self.wb[sheet].cell(row=row + 1, column=col + 1).value
        return default if v is None else v

    def __bool__(self):
        return self.wb is not None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "FX_Hedging_Tool.xlsm"))
    ap.add_argument("--prefill", default=None)
    a = ap.parse_args()
    C = Cache(a.prefill)

    vba_bin = build_vba(os.path.join(ROOT, "build"), [f"Sheet{i + 1}" for i in range(len(SHEETS))],
                        WORKBOOK_CODE)
    wb = xlsxwriter.Workbook(a.out)
    wb.set_vba_name("ThisWorkbook")
    wb.add_vba_project(vba_bin)
    wb.set_properties({"title": "FX Hedging Tool", "subject": "Corporate FX hedging proposals",
                       "author": "Leo Yakass", "comments": "Illustrative market data"})
    ws = {}
    for i, name in enumerate(SHEETS):
        ws[name] = wb.add_worksheet(name)
        ws[name].set_vba_name(f"Sheet{i + 1}")
        ws[name].hide_gridlines(2)

    NAVY = "#14325c"
    IN = {"font_color": "#0000ff", "bg_color": "#fff6d5", "border": 1, "border_color": "#d9c27a"}
    CALC = {"bg_color": "#eef3fa", "border": 1, "border_color": "#c9d6ea"}
    MC = {"bg_color": "#f3f3f1", "border": 1, "border_color": "#d6d5d0"}
    B = {"border": 1, "border_color": "#c9d6ea"}
    F = lambda **k: wb.add_format(k)
    f = {
        "title": F(bold=True, font_size=18, font_color=NAVY),
        "sub": F(italic=True, font_color="#52514e"),
        "section": F(bold=True, font_color="white", bg_color=NAVY),
        "label": F(font_color="#0b0b0b"),
        "in": F(**IN, align="left"), "in4": F(**IN, num_format="0.0000"), "in0": F(**IN, num_format="#,##0"),
        "inpct": F(**IN, num_format="0%"), "in2": F(**IN, num_format="0.00"),
        "calc4": F(**CALC, num_format="0.0000"), "calc0": F(**CALC, num_format="#,##0"),
        "calcpct": F(**CALC, num_format="0%"),
        "mc4": F(**MC, num_format="0.0000"), "mc0": F(**MC, num_format="#,##0"), "mcpct": F(**MC, num_format="0%"),
        "mc2": F(**MC, num_format="0.00"),
        "hdr": F(bold=True, font_color="white", bg_color=NAVY, text_wrap=True, valign="vcenter",
                 align="center", border=1),
        "hdrmc": F(bold=True, font_color="white", bg_color="#5b6b80", text_wrap=True, valign="vcenter",
                   align="center", border=1),
        "name": F(bold=True, **B), "r4": F(num_format="0.0000", **B), "r0": F(num_format="#,##0", **B),
        "status": F(italic=True, font_color="#52514e"),
        "wrap": F(text_wrap=True, valign="top"),
        "big": F(bold=True, font_size=12, text_wrap=True, valign="top"),
        "h2": F(bold=True, font_size=13, font_color=NAVY),
        "s4": F(num_format="0.0000"), "s2": F(num_format="0.00"), "spct": F(num_format="0.0%"),
        "s0": F(num_format="0"), "s6": F(num_format="0.000000"), "sc0": F(num_format="#,##0"),
        "tot0": F(num_format="#,##0", bold=True, top=1),
        "grey": F(font_color="#b5b4ae"),
    }

    def formula(sheet, row, col, text, fmt, default=0):
        ws[sheet].write_formula(row, col, text, fmt, C(sheet, row, col, default))

    # ================================================================== Pricer
    p = ws["Pricer"]
    p.set_column("A:A", 2)
    p.set_column("B:B", 36)
    p.set_column("C:C", 25)
    p.set_column("D:D", 3)
    p.set_column("E:E", 34)
    p.set_column("F:N", 13)
    p.write("B2", "FX Hedging Tool  |  EURUSD corporate hedging proposals", f["title"])
    p.write("B3", "Yellow = input. Blue = live formula. Grey = Monte Carlo output (macro). "
                  "Market data is illustrative.", f["sub"])

    def inp(row, label, value, fmt, name):
        p.write(row, 1, label, f["label"])
        p.write(row, 2, value, f[fmt])
        wb.define_name(name, f"=Pricer!$C${row + 1}")

    p.merge_range("B5:C5", "CLIENT", f["section"])
    inp(5, "Client profile", "Exporter (receives USD)", "in", "Side")
    p.data_validation("C6", {"validate": "list", "source": ["Exporter (receives USD)", "Importer (pays USD)"]})
    inp(6, "USD amount per month", 1_000_000, "in0", "Notional")
    inp(7, "Number of monthly flows (1-24)", 12, "in0", "NMonths")
    p.data_validation("C8", {"validate": "integer", "criteria": "between", "minimum": 1, "maximum": NFIX})
    inp(8, "Budget rate (EURUSD)", 1.18, "in4", "Budget")

    p.merge_range("B11:C11", "MARKET", f["section"])
    inp(11, "EURUSD spot", 1.1650, "in4", "Spot")
    inp(12, "Implied vol shift (vol points)", 0, "in2", "VolShift")
    p.write("B14", "Average forward over the flows", f["label"])
    formula("Pricer", 13, 2, "=FX_AVGFWD(NMonths)", f["calc4"])
    wb.define_name("AvgFwd", "=Pricer!$C$14")
    p.write("B15", "Forward points vs spot (pips)", f["label"])
    formula("Pricer", 14, 2, "=(AvgFwd-Spot)*10000", f["calc0"])

    p.merge_range("B17:C17", "STRUCTURES", f["section"])
    inp(17, "Protection strike (option and collar)", 1.18, "in4", "Protect")
    inp(18, "Participation (participating forward)", 0.5, "inpct", "Participation")
    inp(19, "TARF target (pips of accrued gain)", 1000, "in0", "TarfTarget")
    inp(20, "TARF leverage", 2, "in2", "Leverage")
    p.data_validation("C21", {"validate": "list", "source": [1, 1.5, 2]})
    inp(21, "TARF knock-out type", "full", "in", "KOType")
    p.data_validation("C22", {"validate": "list", "source": ["full", "exact"]})
    inp(22, "Sales margin (bp of notional)", 0, "in0", "MarginBp")

    p.merge_range("B25:C25", "SIMULATION", f["section"])
    inp(25, "Monte Carlo paths", 20000, "in0", "NPaths")
    inp(26, "Random seed", 42, "in0", "Seed")
    inp(27, "Spot drift for statistics", "Spot unchanged (desk view)", "in", "DriftView")
    p.data_validation("C28", {"validate": "list",
                              "source": ["Spot unchanged (desk view)", "Forward (risk-neutral)"]})

    # helper names (formulas, no cell)
    wb.define_name("Phi", '=IF(LEFT(Pricer!$C$6,6)="Export",1,-1)')
    wb.define_name("TotalN", "=Pricer!$C$7*Pricer!$C$8")

    for col, (macro, cap) in zip(["E", "G", "I", "K"], [
            ("PriceAll", "Price all"), ("SolveZeroCost", "Solve zero-cost strikes"),
            ("RunMonteCarlo", "Run Monte Carlo"), ("SetDefaultStrike", "Default protection strike")]):
        p.insert_button(f"{col}5", {"macro": macro, "caption": cap, "width": 170, "height": 34})
    p.merge_range("E8:N8", "", f["status"])
    formula("Pricer", 7, 4, '=IF(Phi=1,"Exporter: a LOWER achieved rate is better (more EUR for each USD). '
                            'Risk: EURUSD rises.","Importer: a HIGHER achieved rate is better (fewer EUR for each '
                            'USD). Risk: EURUSD falls.")', f["status"], "")
    p.merge_range("E10:N10", "Pre-computed results (20,000 paths, seed 42). Change the inputs and click "
                             "'Price all' to re-run." if C else "Click 'Price all' (enable macros).", f["status"])
    wb.define_name("Status", "=Pricer!$E$10")

    heads = ["Strategy", "Strike / rate", "Strike 2 / param", "Mid value to client (EUR)",
             "Upfront premium (EUR)", "Expected rate", "Worst 5%", "Best 5%", "P(miss budget)",
             "P(over-hedged)"]
    p.set_row(11, 30)
    for j, h in enumerate(heads):
        p.write(11, 4 + j, h, f["hdr"] if j < 5 else f["hdrmc"])
    R0 = 12   # first strategy row (0-based) = Excel row 13
    tot = lambda col: f"Schedule!{col}{SROW + NFIX + 2}"   # totals row of the Schedule sheet
    rows = {
        "Unhedged": (None, None, "=0", "=0"),
        "Forward": (f"=SUMPRODUCT(Schedule!$B$3:$B${SROW + NFIX},Schedule!$J$3:$J${SROW + NFIX})/NMonths",
                    None, f"={tot('K')}", "=0"),
        "Vanilla option": ("=Protect", None, f"={tot('P')}", f"={tot('P')}"),
        "Zero-cost collar": ("=Protect", ("in4", 1.1689, "Giveup"), f"={tot('P')}-{tot('U')}", "=0"),
        "Participating forward": (("in4", 1.1926, "PFStrike"), "=Participation", f"={tot('AA')}", "=0"),
        "TARF": (("in4", 1.1245, "TarfStrike"), "=TarfTarget", ("mc0", "TarfPV"), "=0"),
    }
    for i, s in enumerate(STRATS):
        r = R0 + i
        p.write(r, 4, s, f["name"])
        spec = rows[s]
        for j, item in enumerate(spec):
            col = 5 + j
            fmt_num = [f["r4"], f["r4"], f["calc0"], f["calc0"]][j]
            if j == 1 and s == "Participating forward":
                fmt_num = f["calcpct"]
            elif j == 1 and s == "TARF":
                fmt_num = f["calc0"]
            elif j < 2:
                fmt_num = f["calc4"]
            if item is None:
                p.write_blank(r, col, None, f["r4"])
            elif isinstance(item, tuple) and item[0] == "mc0":
                p.write(r, col, C("Pricer", r, col, 0), f["mc0"])
                wb.define_name(item[1], f"=Pricer!${xl_rowcol_to_cell(r, col)[0]}${r + 1}")
            elif isinstance(item, tuple):
                fmt, default, name = item
                p.write(r, col, C("Pricer", r, col, default), f[fmt])
                wb.define_name(name, f"=Pricer!${xl_rowcol_to_cell(r, col).rstrip('0123456789')}${r + 1}")
            else:
                formula("Pricer", r, col, item, fmt_num)
        for j in range(5):
            col = 9 + j
            fmt = [f["mc4"], f["mc4"], f["mc4"], f["mcpct"], f["mcpct"]][j]
            v = C("Pricer", r, col, None)
            if v is None:
                p.write_blank(r, col, None, fmt)
            else:
                p.write(r, col, v, fmt)
    wb.define_name("McTop", "=Pricer!$J$13")
    p.write("E19", "Yellow strikes are solved by 'Solve zero-cost strikes' but can be typed in: the mid value then "
                   "shows what the client pays (negative) or receives.", f["status"])
    p.write("E20", "Mid value: forward, option, collar and participating forward = live Garman-Kohlhagen formulas "
                   "(Schedule sheet). TARF = Monte Carlo.", f["status"])

    p.merge_range("E22:F22", "TARF anatomy (Monte Carlo)", f["section"])
    tl = ["Probability of knock-out", "Expected life (fixings)", "P(at least one leveraged fixing)",
          "Expected number of leveraged fixings"]
    for i, (l, ff) in enumerate(zip(tl, ["mcpct", "mc2", "mcpct", "mc2"])):
        p.write(22 + i, 4, l, f["label"])
        v = C("Pricer", 22 + i, 5, None)
        p.write(22 + i, 5, v, f[ff]) if v is not None else p.write_blank(22 + i, 5, None, f[ff])
    wb.define_name("TarfTop", "=Pricer!$F$23")

    # ================================================================== Schedule (live pricing)
    sc = ws["Schedule"]
    sc.set_column("A:C", 8)
    sc.set_column("D:AA", 11.5)
    sc.freeze_panes(2, 3)
    sc.write("A1", "Live pricing per fixing: forward, discount factors, smile vol and Garman-Kohlhagen "
                   "(rows beyond the number of flows are inactive)", f["h2"])
    groups = [("A", "K", "Market and forward strip"), ("L", "P", "Protection option, strike = Protect"),
              ("Q", "U", "Collar: sold option, strike = Giveup"),
              ("V", "AA", "Participating forward, strike = PFStrike")]
    hdr = ["Fixing", "Active", "T (y)", "Forward", "Fwd pts", "DF USD", "DF EUR", "r USD (cc)", "r EUR (cc)",
           "Fwd hedge rate", "Fwd value (EUR)",
           "Vol", "d1", "d2", "Price (USD/EUR)", "Value (EUR)",
           "Vol", "d1", "d2", "Price (USD/EUR)", "Value (EUR)",
           "Vol", "d1", "d2", "Protect price", "Sold price", "Value (EUR)"]
    sc.set_row(1, 30)
    for j, h in enumerate(hdr):
        sc.write(1, j, h, f["hdr"])
    for g0, g1, txt in groups:
        pass
    PUT = "{K}*F{r}*NORMSDIST(-{d2})-Spot*G{r}*NORMSDIST(-{d1})"
    CALL = "Spot*G{r}*NORMSDIST({d1})-{K}*F{r}*NORMSDIST({d2})"

    def opt(K, d1, d2, r, protect_side=True):
        call, put = CALL.format(K=K, r=r, d1=d1, d2=d2), PUT.format(K=K, r=r, d1=d1, d2=d2)
        cond = "Phi=1" if protect_side else "Phi=-1"
        return f"IF({cond},{call},{put})"

    for i in range(NFIX):
        r0 = SROW + i
        r = r0 + 1
        cells = {
            0: f"=ROW()-{SROW}",
            1: f"=IF(A{r}<=NMonths,1,0)",
            2: f"=A{r}/12",
            3: f"=FX_FORWARD(C{r})",
            4: f"=(D{r}-Spot)*10000",
            5: f"=FX_DFDOM(C{r})",
            6: f"=FX_DFFOR(C{r})",
            7: f"=-LN(F{r})/C{r}",
            8: f"=-LN(G{r})/C{r}",
            9: f"=D{r}-Phi*MarginBp/10000*AvgFwd",
            10: f"=B{r}*Phi*Notional*F{r}*(D{r}/J{r}-1)/Spot",
            11: f"=FX_VOL(Protect,C{r})",
            12: f"=(LN(Spot/Protect)+(H{r}-I{r}+L{r}^2/2)*C{r})/(L{r}*SQRT(C{r}))",
            13: f"=M{r}-L{r}*SQRT(C{r})",
            14: "=" + opt("Protect", f"M{r}", f"N{r}", r, True),
            15: f"=B{r}*Notional/Protect*O{r}/Spot",
            16: f"=FX_VOL(Giveup,C{r})",
            17: f"=(LN(Spot/Giveup)+(H{r}-I{r}+Q{r}^2/2)*C{r})/(Q{r}*SQRT(C{r}))",
            18: f"=R{r}-Q{r}*SQRT(C{r})",
            19: "=" + opt("Giveup", f"R{r}", f"S{r}", r, False),
            20: f"=B{r}*Notional/Giveup*T{r}/Spot",
            21: f"=FX_VOL(PFStrike,C{r})",
            22: f"=(LN(Spot/PFStrike)+(H{r}-I{r}+V{r}^2/2)*C{r})/(V{r}*SQRT(C{r}))",
            23: f"=W{r}-V{r}*SQRT(C{r})",
            24: "=" + opt("PFStrike", f"W{r}", f"X{r}", r, True),
            25: "=" + opt("PFStrike", f"W{r}", f"X{r}", r, False),
            26: f"=B{r}*Notional/PFStrike*(Y{r}-(1-Participation)*Z{r})/Spot",
        }
        fmts = {0: "s0", 1: "s0", 2: "s4", 3: "s4", 4: "s0", 5: "s6", 6: "s6", 7: "spct", 8: "spct",
                9: "s4", 10: "sc0", 11: "spct", 12: "s4", 13: "s4", 14: "s6", 15: "sc0", 16: "spct",
                17: "s4", 18: "s4", 19: "s6", 20: "sc0", 21: "spct", 22: "s4", 23: "s4", 24: "s6",
                25: "s6", 26: "sc0"}
        for col, text in cells.items():
            formula("Schedule", r0, col, text, f[fmts[col]])
    tr = SROW + NFIX + 1          # totals row (0-based) -> Excel row 28
    sc.write(tr, 0, "Total", f["name"])
    for col, letter in [(10, "K"), (15, "P"), (20, "U"), (26, "AA")]:
        formula("Schedule", tr, col, f"=SUM({letter}3:{letter}{SROW + NFIX})", f["tot0"])
    sc.conditional_format(f"A3:AA{SROW + NFIX}", {"type": "formula", "criteria": "=$B3=0", "format": f["grey"]})
    sc.write(tr + 2, 0, "Option premium carried to the average settlement date (EUR):", f["label"])
    formula("Schedule", tr + 2, 9, "=Pricer!I15/FX_DFFOR((1+NMonths)/24)", f["calc0"])
    wb.define_name("PremFV", f"=Schedule!$J${tr + 3}")

    # ================================================================== Scenarios (live)
    s = ws["Scenarios"]
    s.set_column("A:B", 12)
    s.set_column("C:H", 18)
    s.write("A1", "Achieved average rate if EURUSD moves today and stays there (premium included), "
                  "live formulas", f["h2"])
    for j, h in enumerate(["Spot move", "EURUSD"] + STRATS):
        s.write(1, j, h, f["hdr"])
    act = f"Schedule!$B$3:$B${SROW + NFIX}"
    fwdr = f"Schedule!$J$3:$J${SROW + NFIX}"
    for k in range(NSCEN):
        r0 = 2 + k
        r = r0 + 1
        sp = f"B{r}"
        intr = f"(Phi*({sp}-TarfStrike))"
        conv = (f"IF({intr}>0,IF(KOType=\"exact\",MIN(TotalN,Notional*TarfTarget/10000/{intr}),"
                f"Notional*MIN(NMonths,ROUNDUP(TarfTarget/10000/{intr}-1E-12,0))),Leverage*TotalN)")
        cells = {
            0: f"=-0.15+(ROW()-3)*0.005",
            1: f"=Spot*(1+A{r})",
            2: f"={sp}",
            3: f"=NMonths/SUMPRODUCT({act}/{fwdr})",
            4: f"=TotalN/ABS(Phi*TotalN/IF(Phi*({sp}-Protect)>0,Protect,{sp})-PremFV)",
            5: f"=MEDIAN({sp},Protect,Giveup)",
            6: f"=TotalN/ABS(Phi*NMonths*(Notional/{sp}+IF(Phi*({sp}-PFStrike)>0,1,1-Participation)"
               f"*Notional*(1/PFStrike-1/{sp})))",
            7: f"=TotalN/ABS(Phi*(TotalN/{sp}+{conv}*(1/TarfStrike-1/{sp})))",
        }
        for col, text in cells.items():
            formula("Scenarios", r0, col, text, f["spct"] if col == 0 else f["s4"])
    wb.define_name("ScenTop", "=Scenarios!$A$3")

    def axis_range(sheet, r0, r1, c0, c1, step):
        if not C:
            return None
        vals = [C(sheet, r, c, None) for r in range(r0, r1) for c in range(c0, c1)]
        vals = [v for v in vals if isinstance(v, (int, float))]
        return math.floor(min(vals) / step) * step, math.ceil(max(vals) / step) * step

    sx = axis_range("Scenarios", 2, 2 + NSCEN, 1, 2, 0.05) or (0.95, 1.40)
    sy = axis_range("Scenarios", 2, 2 + NSCEN, 2, 8, 0.05) or (0.95, 1.40)
    dy = axis_range("Distribution", 2, 101, 1, 7, 0.02) or (1.0, 1.3)

    ch = wb.add_chart({"type": "scatter", "subtype": "straight"})
    for j, name in enumerate(STRATS):
        ch.add_series({"name": name, "categories": ["Scenarios", 2, 1, 1 + NSCEN, 1],
                       "values": ["Scenarios", 2, 2 + j, 1 + NSCEN, 2 + j],
                       "line": {"color": COLORS[j], "width": 2.25 if j else 1.5,
                                "dash_type": "dash" if j == 0 else "solid"}})
    ch.set_title({"name": "Achieved rate by scenario (EURUSD moves and stays there)",
                  "name_font": {"size": 12, "bold": True}})
    ch.set_x_axis({"name": "EURUSD over the hedging period", "num_format": "0.00", "min": sx[0], "max": sx[1],
                   "major_gridlines": {"visible": True, "line": {"color": "#e6e6e6"}}})
    ch.set_y_axis({"name": "Achieved average rate", "num_format": "0.00", "min": sy[0], "max": sy[1],
                   "major_gridlines": {"visible": True, "line": {"color": "#e6e6e6"}}})
    ch.set_legend({"position": "bottom"})
    ch.set_size({"width": 900, "height": 420})
    p.insert_chart("E28", ch)

    # ================================================================== Distribution (Monte Carlo)
    d = ws["Distribution"]
    d.set_column("A:A", 16)
    d.set_column("B:G", 18)
    d.write("A1", "Monte Carlo outcomes sorted from the client's best to worst (achieved average rate, macro output)",
            f["h2"])
    for j, h in enumerate(["Percentile"] + STRATS):
        d.write(1, j, h, f["hdrmc"])
    for r in range(99):
        for j in range(7):
            v = C("Distribution", 2 + r, j, (r + 1) / 100 if j == 0 else None)
            fmt = f["spct"] if j == 0 else f["s4"]
            d.write(2 + r, j, v, fmt) if v is not None else d.write_blank(2 + r, j, None, fmt)
    wb.define_name("DistTop", "=Distribution!$A$3")
    ch2 = wb.add_chart({"type": "scatter", "subtype": "straight"})
    for j, name in enumerate(STRATS):
        ch2.add_series({"name": name, "categories": ["Distribution", 2, 0, 100, 0],
                        "values": ["Distribution", 2, 1 + j, 100, 1 + j],
                        "line": {"color": COLORS[j], "width": 2.25 if j else 1.5,
                                 "dash_type": "dash" if j == 0 else "solid"}})
    ch2.set_title({"name": "Distribution of outcomes: 1% = best case, 99% = worst case",
                   "name_font": {"size": 12, "bold": True}})
    ch2.set_x_axis({"name": "Percentile of scenarios (best to worst)", "num_format": "0%", "min": 0, "max": 1})
    ch2.set_y_axis({"name": "Achieved average rate", "num_format": "0.00", "min": dy[0], "max": dy[1],
                    "major_gridlines": {"visible": True, "line": {"color": "#e6e6e6"}}})
    ch2.set_legend({"position": "bottom"})
    ch2.set_size({"width": 880, "height": 460})
    d.insert_chart("I2", ch2)

    # ================================================================== Market
    m = ws["Market"]
    m.set_column("A:A", 9)
    m.set_column("B:I", 12)
    m.set_column("J:J", 3)
    m.set_column("K:N", 12)
    m.write("A1", "EURUSD market snapshot (illustrative, edit the yellow cells)", f["h2"])
    m.write("A2", "Rates: continuously compounded zero rates. Smile quoted as ATM, 25-delta risk reversal "
                  "(EUR call minus EUR put) and 25-delta butterfly.", f["sub"])
    mh = ["Tenor", "Years", "USD rate %", "EUR rate %", "ATM vol %", "25d RR %", "25d BF %",
          "Forward", "Fwd points"]
    m.set_row(4, 30)
    for j, h in enumerate(mh):
        m.write(4, j, h, f["hdr"])
    for r in range(15):
        if r < len(MARKET):
            row = MARKET[r]
            m.write(5 + r, 0, row[0], f["in"])
            m.write(5 + r, 1, row[1], f["in4"])
            for j in range(2, 7):
                m.write(5 + r, j, row[j], f["in2"])
            formula("Market", 5 + r, 7, f'=IF(B{6 + r}="","",FX_FORWARD(B{6 + r}))', f["calc4"])
            formula("Market", 5 + r, 8, f'=IF(B{6 + r}="","",FX_POINTS(B{6 + r}))', f["calc0"])
        else:
            for j in range(7):
                m.write_blank(5 + r, j, None, f["in"])
    wb.define_name("MktTable", "=Market!$B$6:$G$20")
    m.write("K4", "Smile check: implied vol (%) by moneyness", f["label"])
    for j, h in enumerate(["Strike / fwd", "1M", "6M", "1Y"]):
        m.write(4, 10 + j, h, f["hdr"])
    for r in range(21):
        m.write(5 + r, 10, -0.06 + 0.006 * r, f["spct"])
        for j, T in enumerate(["1/12", "0.5", "1"]):
            formula("Market", 5 + r, 11 + j, f"=100*FX_VOL(FX_FORWARD({T})*(1+$K{6 + r}),{T})", f["s2"])
    ch3 = wb.add_chart({"type": "scatter", "subtype": "smooth"})
    for j, lab in enumerate(["1M", "6M", "1Y"]):
        ch3.add_series({"name": lab, "categories": ["Market", 5, 10, 25, 10],
                        "values": ["Market", 5, 11 + j, 25, 11 + j],
                        "line": {"color": COLORS[1 + j], "width": 2}})
    ch3.set_title({"name": "Implied volatility smile", "name_font": {"size": 12, "bold": True}})
    ch3.set_x_axis({"name": "strike / forward - 1", "num_format": "0%", "min": -0.06, "max": 0.06})
    ch3.set_y_axis({"name": "implied vol (%)", "num_format": "0.0", "min": 6, "max": 10})
    ch3.set_legend({"position": "bottom"})
    ch3.set_size({"width": 620, "height": 340})
    m.insert_chart("A23", ch3)

    # ================================================================== Pitch (live text)
    pt = ws["Pitch"]
    pt.set_column("A:A", 2)
    pt.set_column("B:B", 110)
    pt.write("B2", "Protecting your USD flows: hedging proposal", f["title"])
    pt.write("B3", "Indicative mid-market levels built from the Pricer sheet. Illustrative market data, not an offer.",
             f["sub"])
    exp_ = 'IF(Phi=1,"receive","pay")'
    lines = [
        ("h2", "Your situation"),
        ("f", f'="You "&{exp_}&" USD "&TEXT(Notional,"#,##0")&" every month for "&NMonths&" months, i.e. USD "'
              f'&TEXT(TotalN/1000000,"0.0")&"m. Your budget rate is "&TEXT(Budget,"0.0000")&'
              f'". EURUSD is "&TEXT(Spot,"0.0000")&" and the average forward over your flows is "&TEXT(AvgFwd,"0.0000")&"."'),
        ("f", '="Without a hedge, our simulations give a "&TEXT(Pricer!M13,"0%")&" chance of missing your budget. '
              'In the worst 5% of scenarios your average rate would be "&TEXT(Pricer!K13,"0.0000")&"."'),
        ("h2", "Our proposals"),
        ("f", '="1. Forward strip: your rate is locked today at "&TEXT(Pricer!F14,"0.0000")&", with no premium '
              'and no upside."'),
        ("f", '="2. Zero-cost collar "&TEXT(Pricer!F16,"0.0000")&" / "&TEXT(Pricer!G16,"0.0000")&": your worst case is '
              '"&TEXT(Pricer!F16,"0.0000")&" and you keep the benefit of a favourable move up to "&TEXT(Pricer!G16,"0.0000")&'
              '". No premium."'),
        ("f", '="3. Vanilla option at "&TEXT(Pricer!F15,"0.0000")&": full upside, for a premium of EUR "'
              '&TEXT(Pricer!I15,"#,##0")&" paid today."'),
        ("f", '="4. Participating forward at "&TEXT(Pricer!F17,"0.0000")&": full protection at that rate, and you keep "'
              '&TEXT(Pricer!G17,"0%")&" of any favourable move. No premium."'),
        ("h2", "A more aggressive alternative, for discussion only"),
        ("f", '="TARF: strike "&TEXT(Pricer!F18,"0.0000")&", "&TEXT(ABS(Pricer!F18-AvgFwd)*10000,"0")&" pips better '
              'than the forward. It stops once you have accumulated "&TEXT(TarfTarget,"#,##0")&" pips of gains (probability "'
              '&TEXT(Pricer!F23,"0%")&"). On an unfavourable fixing you must convert "&Leverage&"x your '
              'flows, which happens at least once in "&TEXT(Pricer!F25,"0%")&" of our scenarios."'),
        ("h2", "Key risks"),
        ("t", "A forward or a collar gives up some or all of the upside. An option premium is paid whatever "
              "happens. A TARF can force conversions above your real flows and leaves you unhedged after "
              "knock-out. Suitability and appropriateness checks apply."),
        ("h2", "Next steps"),
        ("t", "1. Agree the hedge ratio and the share of the budget to protect.  2. ISDA / FBF master agreement "
              "and credit line.  3. Live pricing and execution."),
    ]
    r = 4
    for kind, text in lines:
        if kind == "h2":
            r += 1
            pt.write(r, 1, text, f["h2"])
        elif kind == "f":
            pt.set_row(r, 34)
            formula("Pitch", r, 1, text, f["big"], "")
        else:
            pt.set_row(r, 34)
            pt.write(r, 1, text, f["wrap"])
        r += 1

    # ================================================================== Notes
    n = ws["Notes"]
    n.set_column("A:A", 2)
    n.set_column("B:B", 120)
    notes = [
        ("h2", "How to use"),
        ("t", "1. Enable macros (the workbook recalculates on opening). 2. Fill the yellow cells on the Pricer sheet. "
              "The forward, option, collar and participating forward values, the Scenarios grid and the Pitch "
              "update instantly. 3. 'Solve zero-cost strikes' re-solves the collar give-up, the participating "
              "forward and the TARF strikes. 4. 'Run Monte Carlo' refreshes the TARF value and the statistics. "
              "'Price all' does both."),
        ("h2", "What is a formula and what is VBA"),
        ("t", "Formulas: Schedule sheet (per fixing: forward, discount factors, continuously compounded rates, smile "
              "vol, d1, d2, Garman-Kohlhagen prices with NORMSDIST, option values), Pricer mid values and premium, "
              "Scenarios grid (closed form for a constant spot, TARF included), Pitch text. "
              "VBA: market worksheet functions FX_FORWARD, FX_DFDOM, FX_DFFOR, FX_VOL, FX_POINTS, FX_AVGFWD, "
              "FX_ATMVOL (curve interpolation and smile iteration), bisection solvers, Monte Carlo."),
        ("h2", "Conventions"),
        ("t", "EURUSD is quoted in USD per EUR. The client's home currency is EUR and its flows are in USD. "
              "Exporter (receives USD): risk = EURUSD rises, protection = EUR call / USD put, lower rate is better. "
              "Importer (pays USD): risk = EURUSD falls, protection = EUR put / USD call, higher rate is better."),
        ("h2", "Methodology"),
        ("t", "Forwards: covered interest parity F(T) = S x DF_EUR(T) / DF_USD(T). Smile: implied vol is a quadratic "
              "in forward delta through ATM, 25d RR and 25d BF, solved by fixed-point iteration. Options: "
              "Garman-Kohlhagen with the smile vol of each strike."),
        ("t", "Every structure is a set of conversion legs: n USD converted at a hedge rate R instead of the spot S, "
              "home-currency gain = phi x n x (1/R - 1/S). Mid value = sum DF_USD(t) / S0 x E[gain x S(t)]. "
              "Zero-cost strikes are solved by bisection, minus an optional sales margin."),
        ("t", "TARF: on each fixing, if favourable the client converts N at the strike and accrues the gain; "
              "otherwise it converts leverage x N. It knocks out once the accrued gain reaches the target "
              "(full: last fixing paid in full; exact: scaled to hit the target). Valued by Monte Carlo under "
              "log-normal dynamics with the ATM term structure, common random numbers for the strike solve."),
        ("t", "Statistics: by default the spot is expected to stay where it is (desk view: forward points are a "
              "carry cost, not a forecast). Pricing always uses the risk-neutral forward. Random numbers: "
              "MRG32k3a generator with antithetic variates, so every machine gives the same results."),
        ("h2", "Limitations"),
        ("t", "The TARF ignores the smile (a local-vol model would move the zero-cost strike by tens of pips). "
              "Deterministic rates, no business-day calendar, premium carried at the EUR rate to the average "
              "settlement date. No CVA or capital charge: zero-cost means zero-cost mid-market."),
    ]
    r = 1
    for kind, text in notes:
        if kind == "h2":
            r += 1
            n.write(r, 1, text, f["h2"])
        else:
            n.set_row(r, 62)
            n.write(r, 1, text, f["wrap"])
        r += 1

    # print setup
    for name, sheet in ws.items():
        sheet.set_paper(9)
        sheet.set_portrait() if name == "Pitch" else sheet.set_landscape()
        sheet.fit_to_pages(1, 1 if name in ("Pricer", "Pitch", "Market", "Notes") else 0)
        sheet.set_margins(0.4, 0.4, 0.5, 0.5)
    p.print_area("A1:O50")
    pt.print_area("A1:B30")
    p.activate()
    p.set_first_sheet()
    wb.close()
    print("written", a.out)


if __name__ == "__main__":
    main()
